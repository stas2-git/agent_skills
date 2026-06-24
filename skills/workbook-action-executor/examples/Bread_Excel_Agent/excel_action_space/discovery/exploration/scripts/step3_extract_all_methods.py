"""
Systematic Discovery Script for openpyxl - Step 3
Discovering methods for the key classes we found
"""

import openpyxl
import inspect
import json
from typing import Any

def get_signature_safe(func):
    """Safely get function signature"""
    try:
        sig = inspect.signature(func)
        params = []

        for param_name, param in sig.parameters.items():
            if param_name in ['self', 'cls']:
                continue

            param_info = {
                'name': param_name,
                'required': param.default == inspect.Parameter.empty
            }

            # Get default value if exists
            if param.default != inspect.Parameter.empty:
                if isinstance(param.default, (str, int, float, bool, type(None))):
                    param_info['default'] = param.default
                else:
                    param_info['default'] = str(param.default)

            # Get type annotation if exists
            if param.annotation != inspect.Parameter.empty:
                if hasattr(param.annotation, '__name__'):
                    param_info['type'] = param.annotation.__name__
                else:
                    param_info['type'] = str(param.annotation)

            params.append(param_info)

        return params
    except Exception as e:
        return None

def discover_class_methods(class_path):
    """Discover all methods and properties of a class"""

    # Import the class using importlib
    import importlib

    parts = class_path.split('.')
    module_path = '.'.join(parts[:-1])
    class_name = parts[-1]

    try:
        # Import the module
        module = importlib.import_module(module_path)
        # Get the class from the module
        cls = getattr(module, class_name)
    except Exception as e:
        return {'error': str(e)}

    # Get all members
    all_members = dir(cls)
    public_members = [m for m in all_members if not m.startswith('_')]

    methods = []
    properties = []
    class_variables = []

    for member_name in public_members:
        try:
            member = getattr(cls, member_name)

            # Check what type of member this is
            if inspect.ismethod(member) or inspect.isfunction(member):
                # This is a method
                method_info = {
                    'name': member_name,
                    'docstring': inspect.getdoc(member) or '',
                    'parameters': get_signature_safe(member)
                }
                methods.append(method_info)

            elif isinstance(member, property):
                # This is a property
                prop_info = {
                    'name': member_name,
                    'docstring': inspect.getdoc(member) or '',
                    'readable': member.fget is not None,
                    'writable': member.fset is not None
                }
                properties.append(prop_info)

            elif not callable(member):
                # This might be a class variable
                class_variables.append({
                    'name': member_name,
                    'value': str(member) if not inspect.isclass(member) else f"<class {member.__name__}>",
                    'type': type(member).__name__
                })

        except Exception as e:
            pass

    return {
        'class': class_path,
        'total_public_members': len(public_members),
        'methods': methods,
        'properties': properties,
        'class_variables': class_variables
    }

def discover_openpyxl_step3():
    """Step 3: Discover methods for key classes"""

    print("=" * 60)
    print("STEP 3: Discovering methods for key classes")
    print("=" * 60)

    # Load the classes we found in Step 2
    with open('step2_submodules.json', 'r') as f:
        step2_data = json.load(f)

    key_classes = step2_data['key_classes_found']

    # Discover ALL classes we found in Step 2
    all_classes = []

    # Add key classes
    for name, info in key_classes.items():
        if info.get('accessible'):
            all_classes.append((name, info['location']))

    # Add ALL chart classes found
    chart_classes = [
        ('AreaChart', 'openpyxl.chart.area_chart.AreaChart'),
        ('AreaChart3D', 'openpyxl.chart.area_chart.AreaChart3D'),
        ('BarChart', 'openpyxl.chart.bar_chart.BarChart'),
        ('BarChart3D', 'openpyxl.chart.bar_chart.BarChart3D'),
        ('BubbleChart', 'openpyxl.chart.bubble_chart.BubbleChart'),
        ('DoughnutChart', 'openpyxl.chart.pie_chart.DoughnutChart'),
        ('LineChart', 'openpyxl.chart.line_chart.LineChart'),
        ('LineChart3D', 'openpyxl.chart.line_chart.LineChart3D'),
        ('PieChart', 'openpyxl.chart.pie_chart.PieChart'),
        ('PieChart3D', 'openpyxl.chart.pie_chart.PieChart3D'),
        ('ProjectedPieChart', 'openpyxl.chart.pie_chart.ProjectedPieChart'),
        ('RadarChart', 'openpyxl.chart.radar_chart.RadarChart'),
        ('ScatterChart', 'openpyxl.chart.scatter_chart.ScatterChart'),
        ('StockChart', 'openpyxl.chart.stock_chart.StockChart'),
        ('SurfaceChart', 'openpyxl.chart.surface_chart.SurfaceChart'),
        ('SurfaceChart3D', 'openpyxl.chart.surface_chart.SurfaceChart3D'),
        ('Reference', 'openpyxl.chart.reference.Reference'),
    ]
    all_classes.extend(chart_classes)

    # Add ALL style classes found
    style_classes = [
        ('PatternFill', 'openpyxl.styles.fills.PatternFill'),
        ('GradientFill', 'openpyxl.styles.fills.GradientFill'),
        ('Color', 'openpyxl.styles.colors.Color'),
        ('Side', 'openpyxl.styles.borders.Side'),
        ('NamedStyle', 'openpyxl.styles.named_styles.NamedStyle'),
        ('Protection', 'openpyxl.styles.protection.Protection'),
    ]
    all_classes.extend(style_classes)

    # Add cell types
    cell_classes = [
        ('MergedCell', 'openpyxl.cell.cell.MergedCell'),
        ('ReadOnlyCell', 'openpyxl.cell.read_only.ReadOnlyCell'),
    ]
    all_classes.extend(cell_classes)

    # Add the load_workbook function and other top-level functions
    print("\n📋 Also discovering top-level functions...")
    print("-" * 40)

    # Discover top-level functions
    function_discoveries = {}

    # Import key functions
    from openpyxl import load_workbook
    from openpyxl.utils import (
        get_column_letter, column_index_from_string,
        coordinate_to_tuple, range_boundaries
    )
    from openpyxl.styles import numbers

    functions_to_discover = [
        ('load_workbook', load_workbook),
        ('get_column_letter', get_column_letter),
        ('column_index_from_string', column_index_from_string),
        ('coordinate_to_tuple', coordinate_to_tuple),
        ('range_boundaries', range_boundaries),
    ]

    for func_name, func in functions_to_discover:
        try:
            func_info = {
                'name': func_name,
                'docstring': inspect.getdoc(func) or '',
                'parameters': get_signature_safe(func),
                'is_function': True
            }
            function_discoveries[func_name] = func_info
            print(f"  ✓ Found {func_name} function")
        except Exception as e:
            print(f"  ✗ Could not discover {func_name}: {e}")

    priority_classes = all_classes

    class_discoveries = {}

    for class_nickname, class_path in priority_classes:
        print(f"\n📋 Discovering {class_nickname}...")
        print("-" * 40)

        discovery = discover_class_methods(class_path)

        if 'error' in discovery:
            print(f"  ❌ Error: {discovery['error']}")
        else:
            print(f"  ✓ Found {len(discovery['methods'])} methods")
            print(f"  ✓ Found {len(discovery['properties'])} properties")
            print(f"  ✓ Found {len(discovery['class_variables'])} class variables")

            # Show first few methods as examples
            if discovery['methods']:
                print(f"\n  Sample methods:")
                for method in discovery['methods'][:5]:
                    params = method.get('parameters')
                    if params:
                        param_str = ', '.join([p['name'] for p in params])
                        print(f"    - {method['name']}({param_str})")
                    else:
                        print(f"    - {method['name']}()")

                if len(discovery['methods']) > 5:
                    print(f"    ... and {len(discovery['methods']) - 5} more")

        class_discoveries[class_nickname] = discovery

    # Add function discoveries to the output
    if function_discoveries:
        class_discoveries['_functions'] = function_discoveries

    # Save Step 3 results
    with open('step3_class_methods.json', 'w') as f:
        json.dump(class_discoveries, f, indent=2)

    print(f"\n✅ Step 3 results saved to step3_class_methods.json")

    # Summary
    print("\n" + "=" * 60)
    print("SUMMARY OF DISCOVERED METHODS")
    print("=" * 60)

    total_methods = 0
    for class_name, discovery in class_discoveries.items():
        if 'methods' in discovery:
            method_count = len(discovery['methods'])
            total_methods += method_count
            print(f"{class_name}: {method_count} methods")

    print(f"\nTotal methods discovered: {total_methods}")

    return class_discoveries

if __name__ == "__main__":
    discover_openpyxl_step3()