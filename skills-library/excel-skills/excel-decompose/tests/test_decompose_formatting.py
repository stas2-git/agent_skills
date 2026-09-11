from __future__ import annotations

import importlib.util
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "scripts" / "decompose_excel.py"
spec = importlib.util.spec_from_file_location("decompose_excel", SCRIPT_PATH)
decompose_excel = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(decompose_excel)


def build_rater_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rater"

    ws["B6"] = "Policy Limit"
    ws["C6"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
    ws["C6"].protection = Protection(locked=False)

    ws["D6"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
    ws["D6"].protection = Protection(locked=False)

    ws["B7"] = "Retention"
    border_side = Side(style="thin", color="FF808080")
    ws["C7"].border = Border(
        left=border_side,
        right=border_side,
        top=border_side,
        bottom=border_side,
    )

    ws["B8"] = "Exposure"
    ws["C8"].protection = Protection(locked=False)

    ws["B9"] = "Rate"
    ws["C9"].number_format = "0.0000%"

    ws["B10"] = "Carrier"
    carrier_validation = DataValidation(type="list", formula1='"Admitted,Non-Admitted"', allow_blank=True)
    ws.add_data_validation(carrier_validation)
    carrier_validation.add(ws["C10"])

    ws["B11"] = "Premium"
    ws["C11"] = 12500
    ws["C11"].number_format = '$#,##0.00;[Red]($#,##0.00);"-"'

    ws["B12"] = "Total Premium"
    ws["C12"] = "=C11*1.1"
    ws["C12"].fill = PatternFill(fill_type="solid", fgColor="FFD9EAF7")
    ws["C12"].font = Font(bold=True, color="FF1F4E78")
    ws["C12"].alignment = Alignment(horizontal="right", wrap_text=True)
    ws["C12"].border = Border(top=Side(style="medium", color="FF1F4E78"))

    ws["B13"] = "Attachment"
    decimal_validation = DataValidation(
        type="decimal",
        operator="between",
        formula1=0,
        formula2=100000000,
        allow_blank=True,
        promptTitle="Attachment point",
        prompt="Enter an attachment point.",
    )
    ws.add_data_validation(decimal_validation)
    decimal_validation.add("C13:C20")

    wb.save(path)


def render(path: Path) -> str:
    wb_formula = load_workbook(path, data_only=False)
    wb_values = load_workbook(path, data_only=True)
    try:
        return decompose_excel.render_output(path, wb_formula, wb_values)
    finally:
        wb_formula.close()
        wb_values.close()


def parse_styles(text: str) -> dict[str, dict]:
    styles: dict[str, dict] = {}
    in_styles = False
    for line in text.splitlines():
        if line == "STYLES":
            in_styles = True
            continue
        if in_styles and line == "":
            break
        if in_styles and line.startswith("  s"):
            style_id, payload = line.strip().split(": ", 1)
            styles[style_id] = json.loads(payload)
    return styles


def test_meaningful_blank_cells_styles_and_validations(tmp_path: Path) -> None:
    workbook_path = tmp_path / "rater.xlsx"
    build_rater_workbook(workbook_path)

    text = render(workbook_path)
    styles = parse_styles(text)

    assert "- B6: value=Policy Limit" in text
    assert "- C6: blank | style=" in text
    assert "- C7: blank | style=" in text
    assert "- C8: blank | style=" in text
    assert "- C9: blank | style=" in text
    assert "- C10:" not in text
    assert "- D20:" not in text

    assert text.count('"range":"C10"') == 1
    assert '"type":"list"' in text
    assert '"formula1":"\\"Admitted,Non-Admitted\\""' in text
    assert '"range":"C13:C20"' in text
    assert '"type":"decimal"' in text
    assert '"operator":"between"' in text
    assert '"prompt":"Enter an attachment point."' in text

    assert len(styles) < text.count("| style=s")
    assert any(
        style.get("fill", {}).get("fg_color") == {"type": "rgb", "rgb": "FFFFFF00"}
        for style in styles.values()
    )
    assert any(
        style.get("protection") == {"locked": False}
        for style in styles.values()
    )
    assert any(
        style.get("number_format") == "0.0000%"
        for style in styles.values()
    )
    assert any(
        all(side in style.get("border", {}) for side in ("left", "right", "top", "bottom"))
        and style["border"]["left"]["style"] == "thin"
        and style["border"]["left"]["color"] == {"type": "rgb", "rgb": "FF808080"}
        for style in styles.values()
    )

    assert "- C11: value=12500 | style=" in text
    assert any(
        style.get("number_format") == '$#,##0.00;[Red]($#,##0.00);"-"'
        for style in styles.values()
    )
    assert "- C12: formula=C11*1.1 | style=" in text


def test_style_ids_are_deterministic(tmp_path: Path) -> None:
    workbook_path = tmp_path / "rater.xlsx"
    build_rater_workbook(workbook_path)

    first = parse_styles(render(workbook_path))
    second = parse_styles(render(workbook_path))

    assert first == second
    assert list(first) == [f"s{idx}" for idx in range(1, len(first) + 1)]


def test_formula_cached_value_output_is_preserved(tmp_path: Path) -> None:
    workbook_path = tmp_path / "rater.xlsx"
    build_rater_workbook(workbook_path)

    wb_formula = load_workbook(workbook_path, data_only=False)
    wb_values = load_workbook(workbook_path, data_only=True)
    try:
        wb_values["Rater"]["C12"] = 13750
        text = decompose_excel.render_output(workbook_path, wb_formula, wb_values)
    finally:
        wb_formula.close()
        wb_values.close()

    assert "- C12: formula=C11*1.1 | value=13750 | style=" in text
